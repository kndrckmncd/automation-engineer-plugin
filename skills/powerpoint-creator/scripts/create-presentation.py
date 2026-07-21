#!/usr/bin/env python3
"""
Generate a PowerPoint (.pptx) presentation using the standard Manulife template.

Usage:
    python create-presentation.py --slides-json FILE [--output PATH] [--template PATH]

Inputs:
    --slides-json FILE   Path to JSON file describing the presentation (required)
    --output PATH        Output .pptx file path
                         Default: C:\\Users\\manisea\\OneDrive - Manulife\\Documents\\COPILOT_TEST\\PPT Creation\\<title>.pptx
    --template PATH      Path to the .pptx template file
                         Default: ...\\PPT Creation\\Template.pptx

JSON format:
    {
      "title": "Presentation Title",
      "subtitle": "Optional subtitle",
      "author": "Author Name",
      "slides": [
        {
          "title": "Slide Title",
          "content": ["Bullet 1", "Bullet 2"],
          "notes": "Optional speaker notes"
        }
      ]
    }

Outputs:
    Saves a .pptx file to the output path.
    Prints a JSON summary to stdout on success.

Exit codes:
    0 — success
    1 — input error (missing/invalid JSON, bad path)
    2 — generation error (pptx library failure)
"""

import argparse
import json
import os
import re
import sys
from pathlib import Path
from typing import Any

try:
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.oxml.ns import qn
except ImportError:
    print("Error: python-pptx is not installed. Run: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

DEFAULT_TEMPLATE = (
    r"C:\Users\manisea\OneDrive - Manulife\Documents\COPILOT_TEST\PPT Creation\Template.pptx"
)
DEFAULT_OUTPUT_DIR = (
    r"C:\Users\manisea\OneDrive - Manulife\Documents\COPILOT_TEST\PPT Creation"
)

# Layout indices in the Manulife template
LAYOUT_TITLE   = 0   # "Title Slide"          — idx 0: title, idx 1: subtitle
LAYOUT_CONTENT = 1   # "Title and Content"    — idx 0: title, idx 1: content body
LAYOUT_CHAPTER = 7   # "Chapter"              — idx 0: title only (section divider)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate a PowerPoint from a JSON config.")
    parser.add_argument("--slides-json", required=True, metavar="FILE",
                        help="Path to the slides JSON config file.")
    parser.add_argument("--output", default="", metavar="PATH",
                        help="Output .pptx file path.")
    parser.add_argument("--template", default=DEFAULT_TEMPLATE, metavar="PATH",
                        help="Path to the .pptx template file.")
    return parser.parse_args()


def load_config(path: str) -> dict[str, Any]:
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        print(f"Error: File not found: {path}", file=sys.stderr)
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON in {path}: {e}", file=sys.stderr)
        sys.exit(1)


def safe_filename(title: str) -> str:
    return re.sub(r"[^\w\s\-]", "", title).strip().replace(" ", "_")


def clear_slides(prs: Presentation) -> None:
    """Remove all existing slides from the presentation, keeping layouts and master."""
    xml_slides = prs.slides._sldIdLst
    slide_list = list(prs.slides)
    for slide in slide_list:
        rId = prs.slides._sldIdLst[0].get(qn("r:id"))
        # Remove from the package relationship
        prs.part.drop_rel(rId)
        xml_slides.remove(xml_slides[0])


def set_placeholder_text(slide, idx: int, text: str) -> bool:
    """Set text on a placeholder by idx. Returns True if found."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text = text
            return True
    return False


def set_placeholder_bullets(slide, idx: int, bullets: list[str]) -> bool:
    """Set bullet list on a placeholder by idx. Returns True if found."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            tf = ph.text_frame
            tf.clear()
            for i, bullet in enumerate(bullets):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text  = bullet
                p.level = 0
                p.font.size = Pt(16)
            return True
    return False


def add_title_slide(prs: Presentation, title: str, subtitle: str, author: str) -> None:
    layout = prs.slide_layouts[LAYOUT_TITLE]
    slide  = prs.slides.add_slide(layout)
    set_placeholder_text(slide, 0, title)
    set_placeholder_text(slide, 1, subtitle or author)

    if (prs.slides.notes_slide if hasattr(prs.slides, "notes_slide") else None):
        pass  # no notes on title slide


def add_content_slide(prs: Presentation, slide_title: str,
                      bullets: list[str], notes: str = "") -> None:
    layout = prs.slide_layouts[LAYOUT_CONTENT]
    slide  = prs.slides.add_slide(layout)
    set_placeholder_text(slide, 0, slide_title)
    set_placeholder_bullets(slide, 1, bullets)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def build_presentation(config: dict[str, Any], template_path: str, output_path: str) -> int:
    if not Path(template_path).exists():
        print(f"Error: Template not found: {template_path}", file=sys.stderr)
        sys.exit(1)

    prs = Presentation(template_path)
    clear_slides(prs)

    title    = config.get("title", "Untitled Presentation")
    subtitle = config.get("subtitle", "")
    author   = config.get("author", "")
    slides   = config.get("slides", [])

    add_title_slide(prs, title, subtitle, author)

    for slide in slides:
        add_content_slide(
            prs,
            slide_title = slide.get("title", ""),
            bullets     = slide.get("content", []),
            notes       = slide.get("notes", ""),
        )

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return len(slides) + 1  # +1 for title slide


def main() -> None:
    args   = parse_args()
    config = load_config(args.slides_json)

    title = config.get("title", "Presentation")
    if args.output:
        output_path = args.output
    else:
        output_path = str(Path(DEFAULT_OUTPUT_DIR) / f"{safe_filename(title)}.pptx")

    try:
        total_slides = build_presentation(config, args.template, output_path)
    except Exception as e:
        print(f"Error generating presentation: {e}", file=sys.stderr)
        sys.exit(2)

    print(json.dumps({
        "title":        title,
        "template":     args.template,
        "output":       output_path,
        "total_slides": total_slides,
        "status":       "saved",
    }, indent=2))
    print(f"\nSaved: {output_path}", file=sys.stderr)


if __name__ == "__main__":
    main()




