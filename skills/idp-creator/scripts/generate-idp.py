#!/usr/bin/env python3
"""
Generate a formatted Individual Development Plan (.docx) from a JSON config.

Usage:
    python generate-idp.py --config-json PATH [--output PATH]

Inputs:
    --config-json  Path to JSON file with employee info, goals, development areas
    --output       Output .docx file path (optional; defaults to ~/Downloads/<name>_IDP.docx)

Outputs:
    Formatted .docx file + JSON summary to stdout

Exit codes:
    0 — success
    1 — input error (missing file, invalid JSON)
    2 — generation error (docx write failure)
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Error: python-docx is required. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


# ── Colors ────────────────────────────────────────────────────────────────────
MANULIFE_GREEN  = RGBColor(0x00, 0x6B, 0x3C)
HEADER_DARK     = RGBColor(0x1F, 0x4E, 0x79)
LIGHT_BLUE      = RGBColor(0xBD, 0xD7, 0xEE)
LIGHT_GREEN     = RGBColor(0xE2, 0xEF, 0xDA)
LIGHT_GREY      = RGBColor(0xF2, 0xF2, 0xF2)
WHITE           = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT       = RGBColor(0x26, 0x26, 0x26)

CATEGORY_COLORS = {
    "Technical":   RGBColor(0xDA, 0xE8, 0xFC),
    "Process":     RGBColor(0xD5, 0xE8, 0xD4),
    "Leadership":  RGBColor(0xFF, 0xF2, 0xCC),
    "Growth":      RGBColor(0xF8, 0xCE, 0xCC),
}

STATUS_COLORS = {
    "Completed":   RGBColor(0xD5, 0xE8, 0xD4),
    "In Progress": RGBColor(0xFF, 0xF2, 0xCC),
    "Not Started": RGBColor(0xF2, 0xF2, 0xF2),
    "On Hold":     RGBColor(0xF8, 0xCE, 0xCC),
    "At Risk":     RGBColor(0xF8, 0xCE, 0xCC),
}


def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def bold_run(para, text: str, size: int = 11, color: RGBColor = DARK_TEXT):
    run = para.add_run(text)
    run.bold  = True
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    return run


def normal_run(para, text: str, size: int = 10, color: RGBColor = DARK_TEXT):
    run = para.add_run(text)
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    return run


def set_col_width(table, col_idx: int, width_cm: float):
    for row in table.rows:
        row.cells[col_idx].width = Cm(width_cm)


def add_header_row(table, headers: list[str], bg: RGBColor = HEADER_DARK):
    row = table.rows[0]
    for i, h in enumerate(headers):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.size  = Pt(10)
        run.font.color.rgb = WHITE


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate an IDP Word document.")
    parser.add_argument("--config-json", required=True, help="Path to IDP config JSON file")
    parser.add_argument("--output",      default=None,  help="Output .docx file path")
    return parser.parse_args()


def load_config(path: str) -> dict:
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        print(f"Error: Config file not found: {path}", file=sys.stderr)
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON in {path}: {e}", file=sys.stderr)
        sys.exit(1)


def build_document(cfg: dict) -> Document:
    doc  = Document()
    emp  = cfg.get("employee", {})
    goals = cfg.get("goals", [])
    dev_areas = cfg.get("development_areas", [])

    # ── Page margins ──────────────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin    = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin   = Cm(2.0)
        section.right_margin  = Cm(2.0)

    # ── Title ─────────────────────────────────────────────────────────────────
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run  = title_para.add_run("INDIVIDUAL DEVELOPMENT PLAN")
    title_run.bold = True
    title_run.font.size  = Pt(18)
    title_run.font.color.rgb = MANULIFE_GREEN

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run  = sub_para.add_run(f"Performance Period: {emp.get('period', '2026')}")
    sub_run.font.size  = Pt(11)
    sub_run.font.color.rgb = HEADER_DARK
    sub_run.bold = True

    doc.add_paragraph()

    # ── Employee Info table ───────────────────────────────────────────────────
    info_table = doc.add_table(rows=3, cols=4)
    info_table.style = "Table Grid"
    info_data = [
        [("Employee Name:",  emp.get("name", "")),        ("Position/Title:", emp.get("position", ""))],
        [("Employee Number:", emp.get("employee_number", "")), ("Manager's Name:", emp.get("manager", ""))],
        [("Date Hired:",     emp.get("date_hired", "")),   ("Date Prepared:",  datetime.now().strftime("%B %d, %Y"))],
    ]
    for r_idx, row_data in enumerate(info_data):
        row = info_table.rows[r_idx]
        for c_idx, (label, value) in enumerate(row_data):
            # label cell (cols 0, 2)
            lbl_cell = row.cells[c_idx * 2]
            set_cell_bg(lbl_cell, LIGHT_BLUE)
            p = lbl_cell.paragraphs[0]
            bold_run(p, label, size=10, color=HEADER_DARK)

            # value cell (cols 1, 3)
            val_cell = row.cells[c_idx * 2 + 1]
            p2 = val_cell.paragraphs[0]
            normal_run(p2, value, size=10)

    doc.add_paragraph()

    # ── Section heading helper ────────────────────────────────────────────────
    def section_heading(text: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after  = Pt(4)
        r = p.add_run(text)
        r.bold = True
        r.font.size  = Pt(13)
        r.font.color.rgb = MANULIFE_GREEN

    # ── Goals section ─────────────────────────────────────────────────────────
    section_heading("DEVELOPMENT GOALS")

    goals_table = doc.add_table(rows=1 + len(goals), cols=6)
    goals_table.style = "Table Grid"
    goals_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    add_header_row(goals_table, [
        "Goal Statement", "Key Deliverables", "Action Steps",
        "Target Date", "Status", "Support Needed"
    ])

    for g_idx, goal in enumerate(goals):
        row    = goals_table.rows[g_idx + 1]
        cat    = goal.get("category", "Technical")
        status = goal.get("status", "Not Started")
        cat_bg = CATEGORY_COLORS.get(cat, LIGHT_GREY)
        sts_bg = STATUS_COLORS.get(status, LIGHT_GREY)

        # Goal Statement
        c0 = row.cells[0]
        set_cell_bg(c0, cat_bg)
        p = c0.paragraphs[0]
        bold_run(p, goal.get("goal_statement", ""), size=10)
        if cat:
            p2 = c0.add_paragraph()
            r2 = p2.add_run(f"[{cat}]")
            r2.italic = True
            r2.font.size  = Pt(8)
            r2.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

        # Key Deliverables
        c1 = row.cells[1]
        normal_run(c1.paragraphs[0], goal.get("key_deliverables", ""), size=10)

        # Action Steps
        c2 = row.cells[2]
        steps = goal.get("target_steps", [])
        for s_idx, step in enumerate(steps):
            if s_idx == 0:
                normal_run(c2.paragraphs[0], f"{s_idx+1}. {step}", size=9)
            else:
                p_step = c2.add_paragraph()
                normal_run(p_step, f"{s_idx+1}. {step}", size=9)

        # Target Date
        c3 = row.cells[3]
        c3.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        normal_run(c3.paragraphs[0], goal.get("target_date", "TBD"), size=10)

        # Status
        c4 = row.cells[4]
        set_cell_bg(c4, sts_bg)
        c4.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        bold_run(c4.paragraphs[0], status, size=10)

        # Support Needed
        c5 = row.cells[5]
        normal_run(c5.paragraphs[0], goal.get("support_needed", ""), size=10)

    # Set column widths
    col_widths = [4.5, 3.5, 5.0, 2.0, 2.0, 3.5]
    for ci, w in enumerate(col_widths):
        set_col_width(goals_table, ci, w)

    doc.add_paragraph()

    # ── Development Areas section ─────────────────────────────────────────────
    if dev_areas:
        section_heading("COMPETENCY DEVELOPMENT AREAS")

        dev_table = doc.add_table(rows=1 + len(dev_areas), cols=4)
        dev_table.style = "Table Grid"

        add_header_row(dev_table, [
            "Competency / Skill Area", "Current Level", "Target Level", "Development Actions"
        ])

        for d_idx, area in enumerate(dev_areas):
            row = dev_table.rows[d_idx + 1]
            if d_idx % 2 == 0:
                for cell in row.cells:
                    set_cell_bg(cell, LIGHT_GREY)

            bold_run(row.cells[0].paragraphs[0], area.get("competency", ""), size=10)
            normal_run(row.cells[1].paragraphs[0], area.get("current_level", ""), size=10)
            normal_run(row.cells[2].paragraphs[0], area.get("target_level", ""), size=10)
            normal_run(row.cells[3].paragraphs[0], area.get("actions", ""), size=10)

        set_col_width(dev_table, 0, 4.5)
        set_col_width(dev_table, 1, 2.5)
        set_col_width(dev_table, 2, 2.5)
        set_col_width(dev_table, 3, 6.0)

        doc.add_paragraph()

    # ── Signature block ───────────────────────────────────────────────────────
    section_heading("ACKNOWLEDGEMENT")

    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.style = "Table Grid"

    for c_idx, label in enumerate(["Employee Signature", "Manager Signature"]):
        cell = sig_table.rows[0].cells[c_idx]
        set_cell_bg(cell, LIGHT_BLUE)
        bold_run(cell.paragraphs[0], label, size=10, color=HEADER_DARK)

    for c_idx, value in enumerate([
        f"{emp.get('name', '')}",
        f"{emp.get('manager', '')}"
    ]):
        cell = sig_table.rows[1].cells[c_idx]
        cell.paragraphs[0].paragraph_format.space_before = Pt(18)
        normal_run(cell.paragraphs[0], f"Signature: ________________     Date: __________", size=10)
        p2 = cell.add_paragraph()
        normal_run(p2, value, size=10)

    # ── Footer note ───────────────────────────────────────────────────────────
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("This IDP is a living document. Review and update quarterly with your manager.")
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    return doc


def main() -> None:
    args = parse_args()
    cfg  = load_config(args.config_json)

    emp  = cfg.get("employee", {})
    name = emp.get("name", "Employee").replace(" ", "_")

    if args.output:
        out_path = Path(args.output)
    else:
        out_path = Path.home() / "Downloads" / f"{name}_IDP.docx"

    out_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = build_document(cfg)
        doc.save(str(out_path))
    except Exception as e:
        print(f"Error generating document: {e}", file=sys.stderr)
        sys.exit(2)

    result = {
        "output":  str(out_path),
        "employee": emp.get("name", ""),
        "goals":   len(cfg.get("goals", [])),
        "dev_areas": len(cfg.get("development_areas", [])),
        "status": "saved"
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
