#!/usr/bin/env python3
"""
Generate a Manulife Performance Review document (.docx) — bullet/paragraph format.
Produces: "The What" section, "The How" section, and Ratings on a new page.

Usage:
    python generate-review.py --config-json PATH [--output PATH]

Inputs:
    --config-json  Path to JSON config file
    --output       Output .docx path (defaults to ~/Downloads/<name>_PerformanceReview.docx)

Outputs:
    Formatted .docx + JSON summary to stdout

Exit codes:
    0 — success
    1 — input error
    2 — generation error
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("Error: python-docx is required. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)

# ── Colors ─────────────────────────────────────────────────────────────────
MANULIFE_GREEN  = RGBColor(0x00, 0x6B, 0x3C)
WHAT_BLUE       = RGBColor(0x1F, 0x4E, 0x79)
HOW_PURPLE      = RGBColor(0x5B, 0x2D, 0x8E)
DARK_TEXT       = RGBColor(0x1A, 0x1A, 0x1A)
MUTED           = RGBColor(0x60, 0x60, 0x60)
WHITE           = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BLUE_BG   = RGBColor(0xBD, 0xD7, 0xEE)
LIGHT_GREY_BG   = RGBColor(0xF2, 0xF2, 0xF2)

RATING_COLORS = {
    "Exceptional":      RGBColor(0x00, 0x6B, 0x3C),
    "Highly Effective": RGBColor(0x2E, 0x75, 0xB6),
    "Effective":        RGBColor(0x70, 0xAD, 0x47),
    "Developing":       RGBColor(0xFF, 0x99, 0x00),
}


# ── Helpers ────────────────────────────────────────────────────────────────

def set_cell_bg(cell, rgb: RGBColor):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    hex_color = f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color)
    tcPr.append(shd)


def add_run(para, text, bold=False, italic=False, size=11, color=DARK_TEXT):
    r = para.add_run(text)
    r.bold   = bold
    r.italic = italic
    r.font.size      = Pt(size)
    r.font.color.rgb = color
    return r


def add_para(doc, text="", bold=False, italic=False, size=11, color=DARK_TEXT,
             align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if text:
        add_run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def heading1(doc, text, color=MANULIFE_GREEN):
    """Top-level section heading (large, colored, underlined)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size      = Pt(16)
    r.font.color.rgb = color
    r.font.underline = True
    return p


def heading2(doc, text, color=DARK_TEXT):
    """Subsection heading (bold, slightly larger)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size      = Pt(12)
    r.font.color.rgb = color
    return p


def bullet(doc, text, level=0, size=11):
    """Add a bullet point paragraph."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent   = Cm(0.5 + level * 0.8)
    p.paragraph_format.space_before  = Pt(1)
    p.paragraph_format.space_after   = Pt(2)
    add_run(p, text, size=size, color=DARK_TEXT)
    return p


def divider(doc):
    """Thin horizontal rule effect via a shaded paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(2)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "AAAAAA")
    shd.append(bot)
    pPr.append(shd)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(docx_break_type())
    return p


def docx_break_type():
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    return br


def add_page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def parse_args():
    parser = argparse.ArgumentParser(description="Generate a Performance Review .docx")
    parser.add_argument("--config-json", required=True)
    parser.add_argument("--output",      default=None)
    return parser.parse_args()


def load_config(path):
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        print(f"Error: Config file not found: {path}", file=sys.stderr)
        sys.exit(1)
    except json.JSONDecodeError as e:
        print(f"Error: Invalid JSON: {e}", file=sys.stderr)
        sys.exit(1)


def build_document(cfg):
    doc = Document()
    emp        = cfg.get("employee", {})
    objectives = cfg.get("objectives", [])
    values     = cfg.get("values", [])
    risk       = cfg.get("risk", {})
    ratings    = cfg.get("ratings", {})

    for sec in doc.sections:
        sec.top_margin    = Cm(2.0)
        sec.bottom_margin = Cm(2.0)
        sec.left_margin   = Cm(2.5)
        sec.right_margin  = Cm(2.5)

    # ── Title block ────────────────────────────────────────────────────────
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(2)
    r = t.add_run("PERFORMANCE REVIEW — SELF ASSESSMENT")
    r.bold = True; r.font.size = Pt(20); r.font.color.rgb = MANULIFE_GREEN

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(2)
    r2 = sub.add_run(f"Review Period: {emp.get('period', '2026')}")
    r2.bold = True; r2.font.size = Pt(12); r2.font.color.rgb = WHAT_BLUE

    divider(doc)

    # Employee info — inline paragraph
    info_para = doc.add_paragraph()
    info_para.paragraph_format.space_before = Pt(6)
    info_para.paragraph_format.space_after  = Pt(2)
    fields = [
        ("Name:", emp.get("name","")),
        ("Position:", emp.get("position","")),
        ("Employee #:", emp.get("employee_number","")),
        ("Manager:", emp.get("manager","")),
        ("Date Hired:", emp.get("date_hired","")),
        ("Prepared:", datetime.now().strftime("%B %d, %Y")),
    ]
    for i, (lbl, val) in enumerate(fields):
        if i > 0:
            add_run(info_para, "   |   ", size=10, color=MUTED)
        add_run(info_para, lbl + " ", bold=True, size=10, color=WHAT_BLUE)
        add_run(info_para, val, size=10, color=DARK_TEXT)

    divider(doc)
    doc.add_paragraph()

    # ══════════════════════════════════════════════════════════════════════
    # PART 1 — THE "WHAT"
    # ══════════════════════════════════════════════════════════════════════
    heading1(doc, 'PART 1 — THE "WHAT"', color=WHAT_BLUE)

    intro_what = doc.add_paragraph()
    intro_what.paragraph_format.space_after = Pt(8)
    add_run(intro_what,
        "The following summarizes performance against key objectives for the review period, "
        "assessed on degree of achievement, results, and business impact. "
        "Goals are written using the SMART framework: Specific, Measurable, Attainable, "
        "Results-oriented, and Time-bound.",
        size=10, italic=True, color=MUTED)

    for obj in objectives:
        heading2(doc, obj.get("objective", ""), color=WHAT_BLUE)

        if obj.get("summary"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_run(p, obj["summary"], size=11, color=DARK_TEXT)

        for r_item in obj.get("results", []):
            bullet(doc, r_item, size=11)

        if obj.get("impact"):
            imp = doc.add_paragraph()
            imp.paragraph_format.space_before = Pt(4)
            imp.paragraph_format.space_after  = Pt(8)
            add_run(imp, "Impact: ", bold=True, size=11, color=WHAT_BLUE)
            add_run(imp, obj["impact"], size=11, color=DARK_TEXT)

        divider(doc)

    # ══════════════════════════════════════════════════════════════════════
    # PART 2 — THE "HOW"
    # ══════════════════════════════════════════════════════════════════════
    doc.add_paragraph()
    heading1(doc, 'PART 2 — THE "HOW"', color=HOW_PURPLE)

    intro_how = doc.add_paragraph()
    intro_how.paragraph_format.space_after = Pt(8)
    add_run(intro_how,
        "The following describes how objectives were delivered, assessed through demonstration of "
        "Manulife's 6 Values, degree of influence, and effective management of risk.",
        size=10, italic=True, color=MUTED)

    for val in values:
        heading2(doc, val.get("value", ""), color=HOW_PURPLE)

        if val.get("tagline"):
            tl = doc.add_paragraph()
            tl.paragraph_format.space_after = Pt(3)
            add_run(tl, val["tagline"], size=10, italic=True, color=MUTED)

        if val.get("summary"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_run(p, val["summary"], size=11, color=DARK_TEXT)

        for proof in val.get("proof_points", []):
            bullet(doc, proof, size=11)

        divider(doc)

    # Risk Management
    doc.add_paragraph()
    heading2(doc, "Risk Management", color=HOW_PURPLE)

    if risk.get("summary"):
        rp = doc.add_paragraph()
        rp.paragraph_format.space_after = Pt(4)
        add_run(rp, risk["summary"], size=11, color=DARK_TEXT)

    for q in risk.get("questions", []):
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(4)
        qp.paragraph_format.space_after  = Pt(2)
        add_run(qp, q.get("question",""), bold=True, size=11, color=HOW_PURPLE)
        bullet(doc, q.get("response",""), size=11)

    divider(doc)

    # ══════════════════════════════════════════════════════════════════════
    # RATINGS — new page
    # ══════════════════════════════════════════════════════════════════════
    add_page_break(doc)

    heading1(doc, "PERFORMANCE RATINGS", color=MANULIFE_GREEN)

    add_para(doc,
        "Per Manulife's Global Performance Ratings framework, two separate ratings are assigned: "
        "one for \"The What\" (objectives and results) and one for \"The How\" (values and risk). "
        "The lower of the two ratings determines compensation decisions.",
        size=10, italic=True, color=MUTED, space_after=12)

    divider(doc)

    # What rating block
    what_rating = ratings.get("what", "Effective")
    heading2(doc, 'The "What" — Objectives Rating', color=WHAT_BLUE)
    wr = doc.add_paragraph()
    wr.paragraph_format.space_before = Pt(4)
    wr.paragraph_format.space_after  = Pt(4)
    add_run(wr, "Self-Rating:  ", bold=True, size=13, color=WHAT_BLUE)
    rc = RATING_COLORS.get(what_rating, RATING_COLORS["Effective"])
    r_badge = wr.add_run(f"  {what_rating}  ")
    r_badge.bold = True
    r_badge.font.size = Pt(13)
    r_badge.font.color.rgb = rc

    if ratings.get("what_justification"):
        wj = doc.add_paragraph()
        wj.paragraph_format.space_after = Pt(8)
        add_run(wj, ratings["what_justification"], size=11, color=DARK_TEXT)

    divider(doc)
    doc.add_paragraph()

    # How rating block
    how_rating = ratings.get("how", "Effective")
    heading2(doc, 'The "How" — Values & Risk Rating', color=HOW_PURPLE)
    hr = doc.add_paragraph()
    hr.paragraph_format.space_before = Pt(4)
    hr.paragraph_format.space_after  = Pt(4)
    add_run(hr, "Self-Rating:  ", bold=True, size=13, color=HOW_PURPLE)
    rc2 = RATING_COLORS.get(how_rating, RATING_COLORS["Effective"])
    r_badge2 = hr.add_run(f"  {how_rating}  ")
    r_badge2.bold = True
    r_badge2.font.size = Pt(13)
    r_badge2.font.color.rgb = rc2

    if ratings.get("how_justification"):
        hj = doc.add_paragraph()
        hj.paragraph_format.space_after = Pt(8)
        add_run(hj, ratings["how_justification"], size=11, color=DARK_TEXT)

    divider(doc)
    doc.add_paragraph()

    # Final rating block
    final_rating = ratings.get("final", "Effective")
    heading2(doc, "Final Rating  (lower of What / How)", color=MANULIFE_GREEN)
    fr_p = doc.add_paragraph()
    fr_p.paragraph_format.space_before = Pt(4)
    fr_p.paragraph_format.space_after  = Pt(4)
    add_run(fr_p, "Final Self-Rating:  ", bold=True, size=14, color=MANULIFE_GREEN)
    rc3 = RATING_COLORS.get(final_rating, RATING_COLORS["Effective"])
    r_badge3 = fr_p.add_run(f"  {final_rating}  ")
    r_badge3.bold = True
    r_badge3.font.size = Pt(14)
    r_badge3.font.color.rgb = rc3

    divider(doc)
    doc.add_paragraph()

    # Manager comments
    heading2(doc, "Manager Comments", color=MANULIFE_GREEN)
    mc = doc.add_paragraph()
    mc.paragraph_format.space_before = Pt(30)
    mc.paragraph_format.space_after  = Pt(30)
    add_run(mc, "(To be completed by manager)", size=10, italic=True, color=MUTED)

    divider(doc)
    doc.add_paragraph()

    # Signatures
    heading2(doc, "Acknowledgement", color=MANULIFE_GREEN)
    for label, name in [("Employee", emp.get("name","")), ("Manager", emp.get("manager",""))]:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_before = Pt(12)
        add_run(sp, f"{label}: ", bold=True, size=11, color=DARK_TEXT)
        add_run(sp, "________________________     Date: __________", size=11, color=DARK_TEXT)
        nl = doc.add_paragraph()
        nl.paragraph_format.space_after = Pt(8)
        add_run(nl, f"         {name}", size=10, italic=True, color=MUTED)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(note,
        "The lower of the What and How ratings will be used to determine compensation decisions.",
        italic=True, size=9, color=MUTED)

    return doc


def main():
    args = parse_args()
    cfg  = load_config(args.config_json)
    emp  = cfg.get("employee", {})
    name = emp.get("name","Employee").replace(" ","_")

    out_path = Path(args.output) if args.output else \
               Path.home() / "Downloads" / f"{name}_PerformanceReview.docx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        doc = build_document(cfg)
        doc.save(str(out_path))
    except Exception as e:
        print(f"Error: {e}", file=sys.stderr)
        sys.exit(2)

    print(json.dumps({
        "output":     str(out_path),
        "employee":   emp.get("name",""),
        "objectives": len(cfg.get("objectives",[])),
        "values":     len(cfg.get("values",[])),
        "status":     "saved"
    }, indent=2))


if __name__ == "__main__":
    main()
